# -*- coding: utf-8 -*-
"""Geração das abas de saída em XLSX - Tel.eng.br Stick Corte v0.11 beta"""

from collections import defaultdict, Counter
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, numbers
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
import os
from datetime import datetime

from .optimizer import bar_metrics, kg_for_bars
from .models import Plan, Parameters

VERSION = "0.11 beta"
APP_NAME = "Tel.eng.br - Stick Corte"

COLOR_BLACK = "1A1A1A"
COLOR_HEADER = "1F1F1F"
HEADER_FILL = PatternFill(start_color=COLOR_HEADER, end_color=COLOR_HEADER, fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TOTAL_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
TOTAL_FONT = Font(bold=True, size=11)

def write_sheet(ws, headers, rows, numeric_cols=None, money_cols=None, percent_cols=None):
    ws.append(headers)
    for c in range(1, len(headers)+1):
        cell = ws.cell(1, c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r_idx, row in enumerate(rows, 2):
        ws.append(row)
        for c_idx in range(1, len(headers)+1):
            cell = ws.cell(r_idx, c_idx)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if numeric_cols and (c_idx-1) in numeric_cols:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0'
            if money_cols and (c_idx-1) in money_cols:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = 'R$ #,##0.00'
            if percent_cols and (c_idx-1) in percent_cols:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '0.00%'
    ws.freeze_panes = "A2"
    try:
        ws.auto_filter.ref = ws.dimensions
    except:
        pass
    for col in range(1, ws.max_column + 1):
        max_len = 10
        for r in range(1, min(ws.max_row+1, 200)):
            v = ws.cell(r, col).value
            if v is not None:
                l = len(str(v))
                if l > max_len:
                    max_len = l
        ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 12), 50)
    thin = Side(border_style="thin", color="D9D9D9")
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in r:
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)

def add_cover_sheet(wb, plan: Plan, params: Parameters, logo_tel_path=None, logo_profile_path=None):
    ws = wb.create_sheet("CAPA", 0)
    ws.sheet_properties.tabColor = "000000"
    try:
        if logo_tel_path and os.path.exists(logo_tel_path):
            img = XLImage(logo_tel_path)
            img.width = 400
            img.height = 120
            ws.add_image(img, "B2")
    except:
        pass
    try:
        if logo_profile_path and os.path.exists(logo_profile_path):
            img2 = XLImage(logo_profile_path)
            img2.width = 280
            img2.height = 180
            ws.add_image(img2, "H2")
    except:
        pass
    ws["B6"] = f"{APP_NAME}"
    ws["B6"].font = Font(bold=True, size=20)
    ws["B7"] = f"Versão {VERSION} - Otimizador de corte LSF para Nodus Cad"
    ws["B7"].font = Font(size=12, italic=True, color="666666")
    ws["B9"] = "AUTOR"
    ws["B9"].font = Font(bold=True, size=12)
    ws["B10"] = "Mateus Neves - TEL.ENG.BR | Projetos de Engenharia"
    ws["B11"] = "(47) 9 8836-1017 | contato@tel.eng.br | https://tel.eng.br"
    total_pieces = len(plan.expanded_pieces)
    total_panels = len(set(p.panel for p in plan.expanded_pieces))
    bars_opt = len(plan.bars)
    bars_base = len(plan.baseline_bars)
    bought_opt = sum(b.length_mm for b in plan.bars)
    bought_base = sum(b.length_mm for b in plan.baseline_bars)
    economy_mm = bought_base - bought_opt
    ws["B13"] = "RESUMO DO PROJETO ATUAL"
    ws["B13"].font = Font(bold=True, size=14)
    data = [
        ["Total de peças", total_pieces],
        ["Total de painéis", total_panels],
        ["", ""],
        ["PLANO ORIGINAL (sem otimizar)", f"{bars_base} barras | {bought_base/1000:.1f} m comprado"],
        ["PLANO OTIMIZADO (Stick Corte)", f"{bars_opt} barras | {bought_opt/1000:.1f} m comprado"],
        ["ECONOMIA", f"{economy_mm/1000:.1f} m | {bars_base - bars_opt} barras a menos"],
        ["", ""],
        ["Data de geração", datetime.now().strftime("%d/%m/%Y %H:%M")],
        ["Kerf", f"{params.kerf_mm} mm"],
        ["Sobra mínima útil", f"{params.useful_scrap_min_mm} mm"],
        ["Modo", params.optimization_mode],
    ]
    for i, (k,v) in enumerate(data, 14):
        ws[f"B{i}"] = k
        ws[f"C{i}"] = v
        if "ECONOMIA" in k:
            ws[f"B{i}"].font = Font(bold=True, color="006100", size=12)
            ws[f"C{i}"].font = Font(bold=True, color="006100", size=12)
            ws[f"B{i}"].fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            ws[f"C{i}"].fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    ws["B26"] = "AVISOS LEGAIS - LEIA ANTES DE CORTAR"
    ws["B26"].font = Font(bold=True, size=12, color="FF0000")
    avisos = [
        "1. NODUS CAD é software proprietário (noduscad.com). Este app NÃO tem vínculo com Nodus.",
        "2. Resultado MERAMENTE SUGESTIVO - não substitui julgamento técnico.",
        "3. OBRIGATÓRIO revisão por Eng. Responsável Técnico com ART/RRT.",
        "4. Autor (Mateus Neves / TEL.ENG.BR) NÃO se responsabiliza por perdas ou falhas.",
        "5. Software COMO ESTÁ, sem garantias. Ver AVISO_LEGAL.txt completo.",
    ]
    for i, av in enumerate(avisos, 27):
        ws[f"B{i}"] = av
        ws[f"B{i}"].font = Font(size=9, color="444444")
        try:
            ws.merge_cells(f"B{i}:M{i}")
        except:
            pass
    ws.column_dimensions["B"].width = 45
    ws.column_dimensions["C"].width = 45

def aggregate_scraps(bars):
    counter = Counter()
    for b in bars:
        key = (b.remainder_mm, b.length_mm, b.profile)
        counter[key] += 1
    result = []
    for (rem, length, profile), qty in sorted(counter.items(), key=lambda x: -x[0][0]):
        result.append(f"{qty}x {rem} mm em {length/1000:g}m ({profile})")
    return result

def aggregate_suggestions_for_panel(suggestions):
    counter = Counter()
    for s in suggestions:
        key = (s.target_panel, s.target_mark, s.target_length_mm)
        counter[key] += 1
    parts = []
    for (panel, mark, length), qty in counter.items():
        if qty > 1:
            parts.append(f"{qty}x painel {panel} peça {mark} ({length} mm)")
        else:
            parts.append(f"painel {panel} peça {mark} ({length} mm)")
    return parts

def generate_output(plan: Plan, params: Parameters, output_path: str):
    wb = Workbook()
    wb.remove(wb.active)
    possible_logo_tel = None
    possible_logo_profile = None
    for cand in ["assets/logo_TEL_ENG_BR.png", "assets/Design_sem_nome.png", "Design_sem_nome.png"]:
        if os.path.exists(cand):
            possible_logo_tel = cand
            break
    for cand in ["assets/logo_profiles.png", "assets/icon_profiles.png"]:
        if os.path.exists(cand):
            possible_logo_profile = cand
            break
    add_cover_sheet(wb, plan, params, possible_logo_tel, possible_logo_profile)
    profiles = sorted(set(b.profile for b in plan.bars) | set(b.profile for b in plan.baseline_bars))
    ws = wb.create_sheet("COMPARATIVO")
    ws.sheet_properties.tabColor = "00B050"
    rows = []
    total_base_mm = 0
    total_opt_mm = 0
    total_base_bars = 0
    total_opt_bars = 0
    total_econ_r = 0
    for profile in profiles:
        bars_opt = [b for b in plan.bars if b.profile == profile]
        bars_base = [b for b in plan.baseline_bars if b.profile == profile]
        b3_opt = sum(1 for b in bars_opt if b.length_mm == 3000)
        b6_opt = sum(1 for b in bars_opt if b.length_mm == 6000)
        b3_base = sum(1 for b in bars_base if b.length_mm == 3000)
        b6_base = sum(1 for b in bars_base if b.length_mm == 6000)
        bought_opt, used_opt, useful_opt, real_opt = bar_metrics(bars_opt, params.useful_scrap_min_mm)
        bought_base = sum(b.length_mm for b in bars_base)
        economy_mm = bought_base - bought_opt
        economy_r = economy_mm/1000 * plan.weight_per_meter.get(profile,0) * params.steel_price_per_kg
        total_base_mm += bought_base
        total_opt_mm += bought_opt
        total_base_bars += len(bars_base)
        total_opt_bars += len(bars_opt)
        total_econ_r += economy_r
        rows.append([profile, len(bars_base), f"{b3_base}x3m + {b6_base}x6m", bought_base, len(bars_opt), f"{b3_opt}x3m + {b6_opt}x6m", bought_opt, economy_mm, len(bars_base)-len(bars_opt), economy_r])
    rows.append(["TOTAL", total_base_bars, "", total_base_mm, total_opt_bars, "", total_opt_mm, total_base_mm - total_opt_mm, total_base_bars - total_opt_bars, total_econ_r])
    write_sheet(ws, ["Perfil","Qtd Barras ORIGINAL","Detalhe ORIGINAL","Comprado ORIGINAL (mm)","Qtd Barras OTIMIZADO","Detalhe OTIMIZADO","Comprado OTIMIZADO (mm)","Economia (mm)","Barras economizadas","Economia (R$)"], rows, numeric_cols=[1,3,4,6,7,8], money_cols=[9])
    last_r = ws.max_row
    for c in range(1, ws.max_column+1):
        ws.cell(last_r, c).font = TOTAL_FONT
        ws.cell(last_r, c).fill = TOTAL_FILL
    ws = wb.create_sheet("Resumo OTIMIZADO")
    rows = []
    for profile in profiles:
        bars = [b for b in plan.bars if b.profile == profile]
        b3 = sum(1 for b in bars if b.length_mm == 3000)
        b6 = sum(1 for b in bars if b.length_mm == 6000)
        bought, used, useful, real = bar_metrics(bars, params.useful_scrap_min_mm)
        kg_buy = sum(b.length_mm/1000 * plan.weight_per_meter.get(profile,0) for b in bars)
        kg_used = sum(sum(p.length_mm for p in b.pieces)/1000 * plan.weight_per_meter.get(profile,0) for b in bars)
        loss_pct = (real / bought * 100) if bought else 0
        rows.append([profile, b3, b6, len(bars), bought, used, useful, real, loss_pct/100, kg_buy, kg_used])
    write_sheet(ws, ["Perfil","Barras 3m","Barras 6m","Total barras","Comprado (mm)","Útil (mm)","Sobra aproveitável (mm)","Perda real (mm)","Perda real %","kg comprado","kg usado"], rows, numeric_cols=[1,2,3,4,5,6,7,9,10], percent_cols=[8])
    ws = wb.create_sheet("Plano ORIGINAL")
    ws.sheet_properties.tabColor = "FF0000"
    rows = []
    for b in sorted(plan.baseline_bars, key=lambda x: (x.profile, x.number)):
        seq = " | ".join(f"{p.mark} x {p.length_mm}" for p in b.pieces)
        panel = b.pieces[0].panel if b.pieces else ""
        rows.append([panel, b.profile, b.number, b.length_mm, seq, sum(p.length_mm for p in b.pieces), b.remainder_mm])
    write_sheet(ws, ["Painel","Perfil","Barra nº","Barra (mm)","Sequência","Soma cortes","Sobra"], rows, numeric_cols=[2,3,5,6])
    ws = wb.create_sheet("Plano OTIMIZADO")
    ws.sheet_properties.tabColor = "00B050"
    rows = []
    for b in sorted(plan.bars, key=lambda x: (x.profile, x.number)):
        seq = " | ".join(f"{p.mark} x {p.length_mm}" for p in b.pieces)
        panel = b.pieces[0].panel if b.pieces else ""
        classification = "aproveitável" if b.remainder_mm >= params.useful_scrap_min_mm else "perda real"
        rows.append([panel, b.profile, b.number, b.length_mm, seq, sum(p.length_mm for p in b.pieces), b.remainder_mm, classification])
    write_sheet(ws, ["Painel","Perfil","Barra nº","Barra (mm)","Sequência","Soma cortes","Sobra","Classificação"], rows, numeric_cols=[2,3,5,6])
    ws = wb.create_sheet("Sobras e sugestões")
    rows = []
    for s in plan.suggestions:
        rows.append([s.source_panel, s.source_bar, s.profile, s.scrap_mm, s.target_panel, s.target_mark, s.target_length_mm, s.text])
    write_sheet(ws, ["Origem - painel","Barra nº","Perfil","Sobra (mm)","Painel destino","Peça destino","Compr. destino (mm)","Sugestão"], rows, numeric_cols=[1,3,6])
    ws = wb.create_sheet("Mensagens por painel")
    ws.sheet_properties.tabColor = "FFC000"
    by_panel = defaultdict(list)
    for b in plan.bars:
        if b.pieces:
            by_panel[b.pieces[0].panel].append(b)
    rows = []
    for panel in sorted(by_panel):
        bars = by_panel[panel]
        scraps_agg = aggregate_scraps([b for b in bars if b.remainder_mm >= params.useful_scrap_min_mm])
        if scraps_agg:
            msg = f"Painel {panel}: {len(bars)} barras | Sobras: " + ", ".join(scraps_agg) + "."
            sug = [s for s in plan.suggestions if s.source_panel == panel]
            if sug:
                sug_agg = aggregate_suggestions_for_panel(sug)
                msg += " | Pode reaproveitar para: " + ", ".join(sug_agg) + "."
        else:
            msg = f"Painel {panel}: {len(bars)} barras - sem sobra aproveitável (>={params.useful_scrap_min_mm}mm)."
        rows.append([panel, msg])
    write_sheet(ws, ["Painel","Mensagem resumida"], rows)
    ws = wb.create_sheet("Peças expandidas")
    alloc = {}
    for b in plan.bars:
        for pos, p in enumerate(b.pieces, 1):
            alloc[p.id] = (b, pos)
    rows = []
    for p in plan.expanded_pieces:
        if p.id in alloc:
            b, pos = alloc[p.id]
            status = "Alocada"
            bar_info = f"Barra {b.number} ({b.length_mm} mm)"
        else:
            max_bar = max(params.bar_lengths.get(p.profile, [6000]))
            if p.length_mm > max_bar:
                status = f"NÃO ALOCADA - excede {max_bar}mm - VERIFICAR RT"
                bar_info = "ACIMA DO COMERCIAL"
            else:
                status = "Não alocada"
                bar_info = ""
        rows.append([p.id, p.panel, p.mark, p.profile, p.function, p.length_mm, p.weight_kg, bar_info, status])
    write_sheet(ws, ["ID","Painel","Marca","Perfil","Função","Corte (mm)","Peso (kg)","Barra","Status"], rows, numeric_cols=[0,5,6])
    ws = wb.create_sheet("Lista de compra")
    ws.sheet_properties.tabColor = "0070C0"
    rows = []
    total_kg = 0
    total_r = 0
    total_bars_opt = len(plan.bars)
    for profile in profiles:
        bars = [b for b in plan.bars if b.profile == profile]
        b3 = sum(1 for b in bars if b.length_mm == 3000)
        b6 = sum(1 for b in bars if b.length_mm == 6000)
        kg = sum(b.length_mm/1000 * plan.weight_per_meter.get(profile,0) for b in bars)
        cost = kg * params.steel_price_per_kg
        total_kg += kg
        total_r += cost
        rows.append([profile, b3, b6, len(bars), kg, cost])
    rows.append(["TOTAL", "", "", total_bars_opt, total_kg, total_r])
    write_sheet(ws, ["Perfil","Barras 3m","Barras 6m","Total barras","kg","R$"], rows, numeric_cols=[1,2,3], money_cols=[5])
    last_r = ws.max_row
    for c in range(1, ws.max_column+1):
        ws.cell(last_r, c).font = TOTAL_FONT
        ws.cell(last_r, c).fill = TOTAL_FILL
    ws = wb.create_sheet("Alertas")
    ws.sheet_properties.tabColor = "FF0000"
    rows = [[a.severity, a.panel, a.profile, a.row, a.message] for a in plan.alerts]
    write_sheet(ws, ["Nível","Painel","Perfil","Linha","Mensagem"], rows)
    ws = wb.create_sheet("Corte sob medida")
    grouped = defaultdict(lambda: {"qty": 0, "panels": [], "length": 0})
    for p in plan.expanded_pieces:
        if p.id in {q.id for b in plan.bars for q in b.pieces}:
            key = (p.profile, p.panel, p.mark, p.length_mm)
            grouped[key]["qty"] += 1
            grouped[key]["panels"].append(p.panel)
            grouped[key]["length"] = p.length_mm
    rows = []
    for (profile, panel, mark, length), data in sorted(grouped.items()):
        kg = data["qty"] * length / 1000 * plan.weight_per_meter.get(profile, 0)
        rows.append([profile, panel, mark, length, data["qty"], kg])
    write_sheet(ws, ["Perfil","Painel","Marca","Comprimento (mm)","Quantidade","Peso total (kg)"], rows, numeric_cols=[3,4,5])
    wb.save(output_path)
