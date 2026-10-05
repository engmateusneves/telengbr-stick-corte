# -*- coding: utf-8 -*-
# Tel.eng.br - Stick Corte v0.12 beta - tickavel + colar no CAD
# aqui to meio sem cafe ainda, mas vamos lá

import os
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from collections import defaultdict, Counter

from .parser import read_input, validate_input, infer_weight_per_meter, find_best_sheet
from openpyxl import load_workbook
from .models import Parameters
from .optimizer import build_plan
from .output import generate_output, generate_cad_text_for_panel

VERSION = "0.12 beta"
APP_NAME = "Tel.eng.br - Stick Corte"

DEFAULT_PROFILES = [
    "G90-1.25",
    "89S-41-0.8",
    "140S-1.2",
    "nodus-rfy-metric-067",
    "nodus-rfy-metric-116",
]

class App:
    def __init__(self, root):
        self.root = root
        self.root.title(f"{APP_NAME} - v{VERSION} - Tickavel CAD")
        self.root.geometry("1020x820")
        try:
            if os.path.exists("icon_profiles.ico"):
                self.root.iconbitmap("icon_profiles.ico")
        except:
            pass

        self.input_path = tk.StringVar()
        self.output_path = tk.StringVar()
        self.kerf = tk.StringVar(value="2")
        self.scrap = tk.StringVar(value="100")
        self.mode = tk.StringVar(value="Global")
        self.price = tk.StringVar(value="25,55")
        self.exact_limit = tk.StringVar(value="22")
        self.profile_vars = {}
        self.panel_vars = {}  # pra tickar painel
        self.pieces = []
        self.plan = None
        self.params = None
        self._build()

    def _build(self):
        frame = ttk.Frame(self.root, padding=10)
        frame.pack(fill="both", expand=True)

        ttk.Label(frame, text=f"{APP_NAME}  |  v{VERSION}  |  Tickavel + Colar no CAD", font=("Arial", 13, "bold")).grid(row=0, column=0, columnspan=3, sticky="w")
        ttk.Label(frame, text="TEL.ENG.BR - Mateus Neves (47) 9 8836-1017", font=("Arial", 8)).grid(row=1, column=0, columnspan=3, sticky="w", pady=(0,6))

        ttk.Label(frame, text="Excel do Nodus (.xlsx):").grid(row=2, column=0, sticky="w")
        ttk.Entry(frame, textvariable=self.input_path, width=70).grid(row=3, column=0, columnspan=2, sticky="ew")
        ttk.Button(frame, text="Selecionar...", command=self.select_input).grid(row=3, column=2, padx=6)

        ttk.Label(frame, text="Saida:").grid(row=4, column=0, sticky="w", pady=(8,0))
        ttk.Entry(frame, textvariable=self.output_path, width=70).grid(row=5, column=0, columnspan=2, sticky="ew")
        ttk.Button(frame, text="Salvar como...", command=self.select_output).grid(row=5, column=2, padx=6)

        params = ttk.LabelFrame(frame, text="Parametros", padding=6)
        params.grid(row=6, column=0, columnspan=3, sticky="ew", pady=8)
        ttk.Label(params, text="Kerf:").grid(row=0, column=0, sticky="w")
        ttk.Entry(params, textvariable=self.kerf, width=6).grid(row=0, column=1, sticky="w")
        ttk.Label(params, text="Sobra min:").grid(row=0, column=2, sticky="w", padx=(10,0))
        ttk.Entry(params, textvariable=self.scrap, width=6).grid(row=0, column=3, sticky="w")
        ttk.Label(params, text="R$/kg:").grid(row=0, column=4, sticky="w", padx=(10,0))
        ttk.Entry(params, textvariable=self.price, width=6).grid(row=0, column=5, sticky="w")
        ttk.Label(params, text="Modo:").grid(row=0, column=6, sticky="w", padx=(10,0))
        ttk.Combobox(params, textvariable=self.mode, values=["Por painel", "Global"], state="readonly", width=10).grid(row=0, column=7, sticky="w")

        # lista de paineis com check - v0.12 novo
        self.panel_frame = ttk.LabelFrame(frame, text="Paineis - marca quais quer gerar (tickavel) - v0.12", padding=6)
        self.panel_frame.grid(row=7, column=0, columnspan=3, sticky="nsew", pady=6)
        self.panel_canvas = tk.Canvas(self.panel_frame, height=180)
        self.panel_scroll = ttk.Scrollbar(self.panel_frame, orient="vertical", command=self.panel_canvas.yview)
        self.panel_inner = ttk.Frame(self.panel_canvas)
        self.panel_inner.bind("<Configure>", lambda e: self.panel_canvas.configure(scrollregion=self.panel_canvas.bbox("all")))
        self.panel_canvas.create_window((0,0), window=self.panel_inner, anchor="nw")
        self.panel_canvas.configure(yscrollcommand=self.panel_scroll.set)
        self.panel_canvas.pack(side="left", fill="both", expand=True)
        self.panel_scroll.pack(side="right", fill="y")
        ttk.Label(self.panel_inner, text="Carrega o arquivo primeiro").pack(anchor="w")

        btns = ttk.Frame(frame)
        btns.grid(row=8, column=0, columnspan=3, sticky="ew", pady=6)
        ttk.Button(btns, text="Carregar paineis", command=self.load_profiles).pack(side="left", padx=4)
        ttk.Button(btns, text="Marcar todos", command=lambda: self.toggle_panels(True)).pack(side="left", padx=4)
        ttk.Button(btns, text="Desmarcar todos", command=lambda: self.toggle_panels(False)).pack(side="left", padx=4)
        ttk.Button(btns, text="GERAR EXCEL", command=self.generate).pack(side="right", padx=4)
        ttk.Button(btns, text="Copiar painel p/ CAD", command=self.copy_panel_cad).pack(side="right", padx=4)

        self.status = tk.StringVar(value="Aguardando - v0.12 beta")
        ttk.Label(frame, textvariable=self.status, wraplength=950, font=("Arial", 8)).grid(row=9, column=0, columnspan=3, sticky="w", pady=4)

        frame.columnconfigure(0, weight=1)
        frame.rowconfigure(7, weight=1)

    def select_input(self):
        path = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])
        if path:
            self.input_path.set(path)
            if not self.output_path.get():
                self.output_path.set(os.path.splitext(path)[0] + "_StickCorte_v012.xlsx")

    def select_output(self):
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
        if path:
            self.output_path.set(path)

    def load_profiles(self):
        if not self.input_path.get():
            messagebox.showwarning("Atenção", "Selecione o Excel primeiro")
            return
        try:
            wb = load_workbook(self.input_path.get(), data_only=True, read_only=True)
            sheet = find_best_sheet(wb)
            pieces, _, _ = read_input(self.input_path.get(), sheet_name=sheet)
            self.pieces = pieces
            
            # conta paineis
            by_panel = Counter(p.panel for p in pieces)
            for child in self.panel_inner.winfo_children():
                child.destroy()
            self.panel_vars = {}
            for panel, qtd in sorted(by_panel.items()):
                var = tk.BooleanVar(value=True)
                self.panel_vars[panel] = var
                ttk.Checkbutton(self.panel_inner, text=f"{panel} ({qtd} peças)", variable=var).pack(anchor="w")
            
            self.status.set(f"{len(pieces)} peças da aba '{sheet}' | {len(by_panel)} paineis")
        except Exception as exc:
            messagebox.showerror("Erro", str(exc))

    def toggle_panels(self, state):
        for var in self.panel_vars.values():
            var.set(state)

    def _params(self, pieces):
        inferred = infer_weight_per_meter(pieces)
        profiles = sorted(set(p.profile for p in pieces) | set(DEFAULT_PROFILES))
        bar_lengths = {p: [3000, 6000] for p in profiles}
        optimize = {p: True for p in profiles}
        weights = {p: inferred.get(p) for p in profiles}
        return Parameters(
            bar_lengths=bar_lengths,
            kerf_mm=int(float(self.kerf.get().replace(",", "."))),
            useful_scrap_min_mm=int(float(self.scrap.get().replace(",", "."))),
            optimization_mode=self.mode.get(),
            steel_price_per_kg=float(self.price.get().replace(".", "").replace(",", ".")),
            weight_per_meter=weights,
            optimize_profile=optimize,
            exact_piece_limit=int(self.exact_limit.get()),
        )

    def _filtered_pieces(self):
        # filtra só paineis marcados
        if not self.panel_vars:
            return self.pieces
        selected = {p for p, var in self.panel_vars.items() if var.get()}
        return [x for x in self.pieces if x.panel in selected]

    def generate(self):
        if not self.input_path.get():
            messagebox.showwarning("Atenção", "Selecione o Excel")
            return
        output = self.output_path.get() or os.path.splitext(self.input_path.get())[0] + "_StickCorte_v012.xlsx"
        try:
            self.status.set("Processando... v0.12")
            self.root.update_idletasks()
            wb = load_workbook(self.input_path.get(), data_only=True, read_only=True)
            sheet = find_best_sheet(wb)
            pieces, subtotals, line_alerts = read_input(self.input_path.get(), sheet_name=sheet)
            
            # aplica filtro de paineis tickados
            selected_panels = {p for p, var in self.panel_vars.items() if var.get()} if self.panel_vars else None
            if selected_panels:
                pieces = [p for p in pieces if p.panel in selected_panels]
            
            alerts = validate_input(pieces, subtotals, line_alerts)
            params = self._params(pieces)
            plan = build_plan(pieces, params, alerts, subtotals)
            self.plan = plan
            self.params = params
            generate_output(plan, params, output)
            self.status.set(f"Pronto: {len(plan.baseline_bars)} original vs {len(plan.bars)} otimizado | {output}")
            messagebox.showinfo("Pronto", f"ORIGINAL: {len(plan.baseline_bars)} barras\nOTIMIZADO: {len(plan.bars)} barras\n\n{output}")
        except Exception as exc:
            import traceback
            traceback.print_exc()
            messagebox.showerror("Erro", str(exc))

    def copy_panel_cad(self):
        # pega primeiro painel marcado e copia texto pro CAD
        if not self.plan:
            messagebox.showwarning("Atenção", "Gera o excel primeiro")
            return
        selected = [p for p, var in self.panel_vars.items() if var.get()]
        if not selected:
            messagebox.showwarning("Atenção", "Marca pelo menos 1 painel")
            return
        panel = selected[0]  # copia só o primeiro por enquanto
        txt = generate_cad_text_for_panel(self.plan, panel)
        self.root.clipboard_clear()
        self.root.clipboard_append(txt)
        self.status.set(f"Copiado painel {panel} pra area de transferencia - cola no CAD")
        messagebox.showinfo("Copiado", f"Painel {panel} copiado:\n\n{txt[:300]}...")

def main():
    root = tk.Tk()
    App(root)
    root.mainloop()

if __name__ == "__main__":
    main()
