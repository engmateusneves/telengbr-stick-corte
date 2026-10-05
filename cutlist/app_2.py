# -*- coding: utf-8 -*-
"""Tel.eng.br - Stick Corte v0.11 beta"""

import os
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from .parser import read_input, validate_input, infer_weight_per_meter, find_best_sheet
from openpyxl import load_workbook
from .models import Parameters
from .optimizer import build_plan
from .output import generate_output

VERSION = "0.11 beta"
APP_NAME = "Tel.eng.br - Stick Corte"

DEFAULT_PROFILES = [
    "G90-1.25",
    "89S-41-0.8",
    "140S-1.2",
    "nodus-rfy-metric-067",
    "nodus-rfy-metric-116",
]

class App:
    def __init__(self, root):
        self.root = root
        self.root.title(f"{APP_NAME} - v{VERSION}")
        self.root.geometry("820x640")
        # Tenta carregar ícone
        try:
            if os.path.exists("icon_profiles.ico"):
                self.root.iconbitmap("icon_profiles.ico")
            elif os.path.exists("assets/icon_profiles.ico"):
                self.root.iconbitmap("assets/icon_profiles.ico")
        except:
            pass
        self.input_path = tk.StringVar()
        self.output_path = tk.StringVar()
        self.kerf = tk.StringVar(value="2")
        self.scrap = tk.StringVar(value="100")
        self.mode = tk.StringVar(value="Global")
        self.price = tk.StringVar(value="25,55")
        self.exact_limit = tk.StringVar(value="22")
        self.profile_vars = {}
        self._build()

    def _build(self):
        frame = ttk.Frame(self.root, padding=12)
        frame.pack(fill="both", expand=True)

        # Header com logo textual
        header = ttk.Label(frame, text=f"{APP_NAME}  |  v{VERSION}", font=("Arial", 14, "bold"))
        header.grid(row=0, column=0, columnspan=2, sticky="w", pady=(0,4))
        sub = ttk.Label(frame, text="TEL.ENG.BR - Projetos de Engenharia | Mateus Neves (47) 9 8836-1017", font=("Arial", 9))
        sub.grid(row=1, column=0, columnspan=2, sticky="w", pady=(0,10))

        ttk.Label(frame, text="Arquivo Excel do Nodus (.xlsx):").grid(row=2, column=0, sticky="w")
        ttk.Entry(frame, textvariable=self.input_path, width=75).grid(row=3, column=0, sticky="ew")
        ttk.Button(frame, text="Selecionar...", command=self.select_input).grid(row=3, column=1, padx=6)

        ttk.Label(frame, text="Arquivo de saída:").grid(row=4, column=0, sticky="w", pady=(12,0))
        ttk.Entry(frame, textvariable=self.output_path, width=75).grid(row=5, column=0, sticky="ew")
        ttk.Button(frame, text="Salvar como...", command=self.select_output).grid(row=5, column=1, padx=6)

        params = ttk.LabelFrame(frame, text="Parâmetros de corte", padding=10)
        params.grid(row=6, column=0, columnspan=2, sticky="ew", pady=12)
        ttk.Label(params, text="Kerf (mm):").grid(row=0, column=0, sticky="w")
        ttk.Entry(params, textvariable=self.kerf, width=10).grid(row=0, column=1, sticky="w")
        ttk.Label(params, text="Sobra mínima útil (mm):").grid(row=0, column=2, sticky="w", padx=(20,0))
        ttk.Entry(params, textvariable=self.scrap, width=10).grid(row=0, column=3, sticky="w")
        ttk.Label(params, text="Preço do aço (R$/kg):").grid(row=1, column=0, sticky="w", pady=8)
        ttk.Entry(params, textvariable=self.price, width=10).grid(row=1, column=1, sticky="w")
        ttk.Label(params, text="Limite busca exata:").grid(row=1, column=2, sticky="w", padx=(20,0))
        ttk.Entry(params, textvariable=self.exact_limit, width=10).grid(row=1, column=3, sticky="w")
        ttk.Label(params, text="Modo:").grid(row=2, column=0, sticky="w")
        ttk.Combobox(params, textvariable=self.mode, values=["Por painel", "Global"], state="readonly", width=18).grid(row=2, column=1, sticky="w")
        ttk.Label(params, text="Global = melhor aproveitamento | Por painel = melhor rastreio").grid(row=2, column=2, columnspan=2, sticky="w", padx=(20,0))

        self.profile_frame = ttk.LabelFrame(frame, text="Perfis (auto-detectados)", padding=8)
        self.profile_frame.grid(row=7, column=0, columnspan=2, sticky="ew")
        ttk.Label(self.profile_frame, text="Selecione o arquivo e clique em Carregar perfis.").pack(anchor="w")

        btn_frame = ttk.Frame(frame)
        btn_frame.grid(row=8, column=0, columnspan=2, sticky="ew", pady=10)
        ttk.Button(btn_frame, text="Carregar perfis", command=self.load_profiles).pack(side="left")
        ttk.Button(btn_frame, text="GERAR - Stick Corte", command=self.generate).pack(side="right")

        self.status = tk.StringVar(value="Aguardando arquivo. Versão 0.11 beta - Uso deve ser validado por RT.")
        ttk.Label(frame, textvariable=self.status, wraplength=780, font=("Arial", 8)).grid(row=9, column=0, columnspan=2, sticky="w")

        frame.columnconfigure(0, weight=1)

    def select_input(self):
        path = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])
        if path:
            self.input_path.set(path)
            if not self.output_path.get():
                self.output_path.set(os.path.splitext(path)[0] + "_StickCorte_v011.xlsx")

    def select_output(self):
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
        if path:
            self.output_path.set(path)

    def load_profiles(self):
        if not self.input_path.get():
            messagebox.showwarning("Atenção", "Selecione o Excel primeiro.")
            return
        try:
            wb = load_workbook(self.input_path.get(), data_only=True, read_only=True)
            sheet = find_best_sheet(wb)
            pieces, _, _ = read_input(self.input_path.get(), sheet_name=sheet)
            profiles = sorted(set(p.profile for p in pieces) | set(DEFAULT_PROFILES))
            for child in self.profile_frame.winfo_children():
                child.destroy()
            for p in profiles:
                var = tk.BooleanVar(value=True)
                self.profile_vars[p] = var
                ttk.Checkbutton(self.profile_frame, text=p, variable=var).pack(anchor="w")
            self.status.set(f"{len(pieces)} peças lidas da aba '{sheet}'; {len(profiles)} perfis | Pronto para gerar.")
        except Exception as exc:
            messagebox.showerror("Erro", str(exc))

    def _params(self, pieces):
        inferred = infer_weight_per_meter(pieces)
        profiles = sorted(set(p.profile for p in pieces) | set(DEFAULT_PROFILES))
        bar_lengths = {p: [3000, 6000] for p in profiles}
        optimize = {p: self.profile_vars.get(p, tk.BooleanVar(value=True)).get() for p in profiles}
        weights = {p: inferred.get(p) for p in profiles}
        return Parameters(
            bar_lengths=bar_lengths,
            kerf_mm=int(float(self.kerf.get().replace(",", "."))),
            useful_scrap_min_mm=int(float(self.scrap.get().replace(",", "."))),
            optimization_mode=self.mode.get(),
            steel_price_per_kg=float(self.price.get().replace(".", "").replace(",", ".")),
            weight_per_meter=weights,
            optimize_profile=optimize,
            exact_piece_limit=int(self.exact_limit.get()),
        )

    def generate(self):
        if not self.input_path.get():
            messagebox.showwarning("Atenção", "Selecione o Excel.")
            return
        output = self.output_path.get()
        if not output:
            output = os.path.splitext(self.input_path.get())[0] + "_StickCorte_v011.xlsx"
            self.output_path.set(output)
        try:
            self.status.set("Processando... Otimizando corte (pode levar 1-2 min para 600+ peças)...")
            self.root.update_idletasks()
            wb = load_workbook(self.input_path.get(), data_only=True, read_only=True)
            sheet = find_best_sheet(wb)
            pieces, subtotals, line_alerts = read_input(self.input_path.get(), sheet_name=sheet)
            alerts = validate_input(pieces, subtotals, line_alerts)
            params = self._params(pieces)
            plan = build_plan(pieces, params, alerts, subtotals)
            generate_output(plan, params, output)
            self.status.set(f"Concluído: {len(plan.baseline_bars)} barras ORIGINAL vs {len(plan.bars)} OTIMIZADO | Economia {len(plan.baseline_bars)-len(plan.bars)} barras | Saída: {output}")
            messagebox.showinfo("Concluído - Stick Corte v0.11", f"Plano ORIGINAL: {len(plan.baseline_bars)} barras\nPlano OTIMIZADO: {len(plan.bars)} barras\n\nEconomia: {len(plan.baseline_bars)-len(plan.bars)} barras\n\nArquivo:\n{output}\n\nLembre-se: validar com RT!")
        except Exception as exc:
            self.status.set("Erro.")
            import traceback
            traceback.print_exc()
            messagebox.showerror("Erro", str(exc))

def main():
    root = tk.Tk()
    App(root)
    root.mainloop()

if __name__ == "__main__":
    main()
