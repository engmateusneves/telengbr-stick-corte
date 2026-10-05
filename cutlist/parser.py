# -*- coding: utf-8 -*-
# parser do nodus - lê o excel que vem do cad
# o nodus muda o nome da aba direto, então tem que procurar

from collections import defaultdict
from decimal import Decimal, InvalidOperation
from typing import Dict, List, Tuple
import re
from openpyxl import load_workbook

from .models import Piece, Alert

REQUIRED = [
    "Painel", "Marca da peça", "Perfil / material", "Função da peça",
    "Qtd. (peças)", "Corte por peça (mm)", "Comprimento total (m)", "Peso (kg)"
]

def parse_number_ptbr(value, integer=False):
    # converte numero que vem com virgula, ponto, mm
    # as vezes vem "2.500,50 mm" ai tem que limpar
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return int(round(value)) if integer else float(value)

    s = str(value).strip().replace("\xa0", " ")
    s = s.replace("mm", "").replace("m", "").strip()
    if not s:
        return None

    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    else:
        if integer and re.fullmatch(r"-?\d{1,3}(?:\.\d{3})+", s):
            s = s.replace(".", "")

    try:
        n = Decimal(s)
    except InvalidOperation:
        raise ValueError(f"Número inválido: {value!r}")
    return int(round(n)) if integer else float(n)


def normalize_header(v):
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).strip()).lower()


def find_header_row(ws):
    # procura linha que tem "Painel" e "Corte por peça"
    # varre as primeiras 30 linhas
    max_r = ws.max_row or 30
    max_c = ws.max_column or 20
    for r in range(1, min(max_r, 30) + 1):
        values = [normalize_header(ws.cell(r, c).value) for c in range(1, min(max_c, 20) + 1)]
        if "painel" in values and any("corte por peça" in x for x in values):
            return r, {normalize_header(ws.cell(r, c).value): c for c in range(1, min(max_c, 20) + 1)}
    raise ValueError('Cabeçalho não encontrado nesta aba.')


def find_best_sheet(wb):
    # tenta achar a aba que tem cabeçalho do nodus
    # testa todas, a primeira que tiver cabeçalho é a certa
    for name in wb.sheetnames:
        try:
            ws = wb[name]
            find_header_row(ws)
            return name
        except:
            continue
    return wb.sheetnames[0] if wb.sheetnames else None


def parse_panel_marker(value):
    if value is None:
        return None
    m = re.search(r"painel\s*:\s*(.+)", str(value), flags=re.I)
    return m.group(1).strip() if m else None


def parse_subtotal_marker(value):
    if value is None:
        return None
    m = re.search(r"subtotal\s+do\s+painel\s*:\s*(.+)", str(value), flags=re.I)
    return m.group(1).strip() if m else None


def read_input(path: str, sheet_name: str = None):
    wb = load_workbook(path, data_only=True, read_only=False)
    
    # se não passou aba ou a aba não existe, tenta achar sozinho
    if sheet_name is None or sheet_name not in wb.sheetnames:
        best = find_best_sheet(wb)
        if best is None:
            raise ValueError(f'Nenhuma aba válida encontrada. Abas: {wb.sheetnames}')
        sheet_name = best

    if sheet_name not in wb.sheetnames:
        raise ValueError(f'Aba "{sheet_name}" não existe. Abas: {wb.sheetnames}')

    ws = wb[sheet_name]
    try:
        header_row, headers = find_header_row(ws)
    except ValueError:
        best = find_best_sheet(wb)
        if best and best != sheet_name:
            ws = wb[best]
            sheet_name = best
            header_row, headers = find_header_row(ws)
        else:
            raise ValueError(f'Não achei cabeçalho em nenhuma aba. Abas: {wb.sheetnames}')

    def col(name):
        n = normalize_header(name)
        for k, v in headers.items():
            if k == n or n in k:
                return v
        raise ValueError(f'Coluna não encontrada: {name}')

    cols = {name: col(name) for name in REQUIRED}
    pieces: List[Piece] = []
    alerts: List[Alert] = []
    subtotals: Dict[str, dict] = {}
    current_panel = None
    pid = 1

    for r in range(header_row + 1, ws.max_row + 1):
        row_values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        nonempty = [x for x in row_values if x not in (None, "")]
        if not nonempty:
            continue

        first_text = str(nonempty[0]).strip()
        marker = parse_panel_marker(first_text)
        if marker:
            current_panel = marker
            continue

        subtotal = parse_subtotal_marker(first_text)
        if subtotal:
            m = re.search(r"(\d+)\s*peças?", first_text, flags=re.I)
            total_m = re.search(r"([\d\.,]+)\s*m\b", first_text, flags=re.I)
            total_kg = re.search(r"([\d\.,]+)\s*kg\b", first_text, flags=re.I)
            subtotals[subtotal] = {
                "pieces": int(m.group(1)) if m else None,
                "meters": parse_number_ptbr(total_m.group(1)) if total_m else None,
                "kg": parse_number_ptbr(total_kg.group(1)) if total_kg else None,
            }
            current_panel = subtotal
            continue

        panel = ws.cell(r, cols["Painel"]).value or current_panel
        mark = ws.cell(r, cols["Marca da peça"]).value
        profile = ws.cell(r, cols["Perfil / material"]).value
        function = ws.cell(r, cols["Função da peça"]).value
        qty = parse_number_ptbr(ws.cell(r, cols["Qtd. (peças)"]).value, integer=True)
        cut = parse_number_ptbr(ws.cell(r, cols["Corte por peça (mm)"]).value, integer=True)
        total_m = parse_number_ptbr(ws.cell(r, cols["Comprimento total (m)"]).value)
        weight = parse_number_ptbr(ws.cell(r, cols["Peso (kg)"]).value)

        if not mark or not profile or qty is None or cut is None:
            continue
        if qty <= 0:
            continue

        panel = str(panel).strip() if panel is not None else ""
        profile = str(profile).strip()
        function = str(function or "").strip()
        mark = str(mark).strip()

        for i in range(qty):
            pieces.append(Piece(
                id=pid, panel=panel, mark=mark, profile=profile,
                function=function, length_mm=cut,
                weight_kg=(weight / qty if weight is not None else 0.0),
                source_row=r, source_qty=qty
            ))
            pid += 1

    return pieces, subtotals, alerts

def validate_input(pieces, subtotals, line_alerts):
    from collections import defaultdict
    from .models import Alert
    alerts = list(line_alerts)
    by_panel = defaultdict(list)
    for p in pieces:
        by_panel[p.panel].append(p)
    return alerts

def infer_weight_per_meter(pieces):
    from collections import defaultdict
    sums = defaultdict(lambda: [0.0, 0.0])
    for p in pieces:
        sums[p.profile][0] += p.weight_kg
        sums[p.profile][1] += p.length_mm / 1000.0
    return {
        profile: (vals[0] / vals[1] if vals[1] > 0 else 0.0)
        for profile, vals in sums.items()
    }
